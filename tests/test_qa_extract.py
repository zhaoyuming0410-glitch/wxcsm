# -*- coding: utf-8 -*-
"""客户问题答复清单：抽取与导出的回归测试。

重点固化的都是**真实踩过的坑**，不是理论用例：
  1. 「能不能」里的子串「不能」把正常提问判成报错
  2. 「申请单」里的「请」被当成请求动词
  3. 「比我想的问题多」里的裸「想」让一句评论被判成需求
  4. `_PLEASANTRY` 以「那」开头就算客套，误杀「那能不能帮我们看一下配置？」
  5. 「采购流没问题」这类假问题不能被收
  6. XML/app 载荷残渣（"2598…@openim:"、"…</aeskey"）不能进表格
  7. overwrite 不能把同一工作簿里的另一张表删掉

直接运行：python tests/test_qa_extract.py
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook

from core import exporter as E
from core import qa_extract as Q
from core.models import (
    QA_STATUS_DONE, QA_STATUS_PENDING, QA_STATUS_PROMISED, ChatBundle, Contact, Message,
    Summary,
)
from core.summarizer import _strip_message_noise

BOT = "壬癸客服"


def mk(messages):
    """messages: (天, 时, 分, 发送人, 文本) —— 发送人为 BOT 表示我方。"""
    msgs = []
    for d, hh, mm, who, text in messages:
        msgs.append(Message(ts=datetime(2026, 8, d, hh, mm), sender=who, text=text,
                            is_self=(who == BOT)))
    return ChatBundle(contact=Contact(cid="c1", name="测试客户", kind="friend"),
                      start=date(2026, 8, 1), end=date(2026, 8, 31), messages=msgs)


class TestClassify(unittest.TestCase):
    def test_提问_能力询问不能被子串不能误判成报错(self):
        # 坑 1：「能不能」含「不能」
        self.assertEqual(Q.classify("另外续费方案能不能给个优惠，我们今年预算收紧了"), "需求")
        self.assertEqual(Q.classify("那能不能帮我们看一下配置"), "需求")
        self.assertEqual(Q.classify("差旅标准能不能按我们自己的标准分档呢"), "提问")

    def test_真实的不能仍然是报错(self):
        self.assertEqual(Q.classify("我们这边不能上传发票"), "报错")

    def test_申请单里的请不算请求动词(self):
        # 坑 2
        s = "早上好，我们财务提了个问题，差旅申请单里的城市级别能不能按我们自己的标准分档"
        self.assertEqual(Q.classify(s), "提问")

    def test_裸想不能把评论判成需求(self):
        # 坑 3
        self.assertIsNone(Q.classify("看到了，比我想的问题多"))
        # 但"想+动作"仍是需求
        self.assertEqual(Q.classify("我们下个月就想跑起来"), "需求")

    def test_那开头的正经提问不能被当客套杀掉(self):
        # 坑 4
        self.assertIsNotNone(Q.classify("那能不能帮我们看一下配置"))
        self.assertIsNone(Q.classify("那"))            # 真的客套仍是噪音

    def test_采购流没问题不是报错(self):
        # 坑 5
        self.assertIsNone(Q.classify("试了，采购流没问题"))
        self.assertEqual(Q.classify("但审批到第二级卡住了，说没有匹配的审批人"), "报错")

    def test_各类标签(self):
        self.assertEqual(Q.classify("差旅那条第三级审批人显示空的"), "报错")
        self.assertEqual(Q.classify("实名这块我们推不动，有什么办法"), "投诉")
        self.assertEqual(Q.classify("今年用得不太顺"), "投诉")
        self.assertEqual(Q.classify("把「电子发票查验」这块讲细一点"), "需求")
        self.assertEqual(Q.classify("另外我们有个诉求，希望采购和差旅两条审批流分开"), "需求")
        self.assertEqual(Q.classify("人员导入我们提供Excel就行吗"), "提问")
        # 陈述约束也算隐含需求（漏了会丢真实诉求）
        self.assertEqual(
            Q.classify("同步过来的部门要和我们费控预算中心对上，不然预算卡不住"), "需求")


class TestPayloadFilter(unittest.TestCase):
    def test_载荷残渣被识别(self):
        for s in ("10000000000000001:57</aeskey", "10000000000000001:57<",
                  "10000000000000001:", '<?xml version="1.0"?><msg>'):
            self.assertTrue(Q._looks_like_payload(s), s)

    def test_正常文本不被误杀(self):
        for s in ("赵老师:好的", "时间:10点", "收到，会尽快安排时间配置后台的",
                  "这个功能报错了", "2026-09-11", "sso.bingding-tech.com"):
            self.assertFalse(Q._looks_like_payload(s), s)


class TestMessageCleanup(unittest.TestCase):
    def test_发送人标签藏在XML或占位后面也能剥掉(self):
        # 坑 6：标签被 [类型xxx]/XML 挡在前面时，@提及清理会把 @openim 吃掉，
        # 只剩半截 id 留在正文
        self.assertEqual(
            _strip_message_noise("10000000000000001@openim:\n在架构内开通权限就可以使用"),
            "在架构内开通权限就可以使用")
        self.assertEqual(
            _strip_message_noise("wxid_lisi:\n@张三 @李四 有个问题咨询一下"),
            "有个问题咨询一下")

    def test_未闭合CDATA不残留(self):
        self.assertEqual(
            _strip_message_noise("<![CDATA[，我们现在配置的是支付前强制上传"),
            "我们现在配置的是支付前强制上传")

    def test_纯XML消息清洗后为空(self):
        # 只有标记、没有 title/des 正文的消息，清洗后应为空
        self.assertEqual(_strip_message_noise(
            '[类型244813135921] 2598@openim:\n<?xml version="1.0"?>\n'
            '<msg><appmsg><action /><type>57</type><showtype>0</showtype>'
            '</appmsg></msg>'), "")

    def test_企微图文消息的正文要捞出来而不是丢掉(self):
        # 真实坑：企微/OpenIM 桥接号（local_type 244813135921）把真正的对话文字放在
        # <title> 里，<type>57</type> 只是数字标记。过去先跑 XML 标签剥离会把 CDATA
        # 正文整段吃掉、只留 "57"，导致「我方答复」变成一句 57 或"已答复。"。
        # 正文有两种写法，真实数据里都出现过。
        self.assertEqual(
            _strip_message_noise(
                '[类型244813135921] 2598@openim:\n<?xml version="1.0"?>\n'
                '<msg><appmsg><title><![CDATA[@王五 支付宝企业码和美团没有关系]]></title>'
                '<des /><type>57</type><soundtype>0</soundtype></appmsg></msg>'),
            "支付宝企业码和美团没有关系")
        self.assertEqual(
            _strip_message_noise(
                '[类型244813135921] wxid_wangwu:\n<?xml version="1.0"?>\n'
                '<msg><appmsg><title>我看您发的链接是11:00的</title>'
                '<des /><type>57</type><soundtype>0</soundty'),
            "我看您发的链接是11:00的")

    def test_正文里的时间不会被当成发送人标签(self):
        # 中文标签规则放宽会把「我看您发的链接是11:」当成账号名剥掉，只剩「00的」
        for raw, want in (
            ("我看您发的链接是11:00的", "我看您发的链接是11:00的"),
            ("会议时间:2026/07/28", "会议时间:2026/07/28"),
            ("张三:今天发你", "今天发你"),
        ):
            self.assertEqual(_strip_message_noise(raw), want)


class TestExtract(unittest.TestCase):
    def test_配对与状态(self):
        b = mk([
            (5, 9, 12, "张三", "审批到第二级卡住了，说没有匹配的审批人"),
            (5, 9, 20, BOT, "收到，是审批流里海南分组没绑对应审批节点，我今晚修"),
            (6, 10, 0, "张三", "培训能安排吗"),
            (7, 10, 0, "张三", "另外续费方案能不能给个优惠呢"),
        ])
        items = Q.extract_qa(b, {})
        kinds = [it.kind for it in items]
        self.assertIn("报错", kinds)
        first = [it for it in items if it.kind == "报错"][0]
        self.assertEqual(first.status, QA_STATUS_DONE)
        self.assertIn("海南分组", first.answer)
        self.assertIn("客户[08-05", first.raw)          # 原始记录带双方原话
        self.assertIn("我方[08-05", first.raw)
        # 培训那条我方没回 -> 未答复
        pend = [it for it in items if it.kind == "需求" and "培训" in it.question]
        self.assertTrue(pend)
        self.assertEqual(pend[0].status, QA_STATUS_PENDING)
        self.assertEqual(pend[0].answer, "")

    def test_同一条消息多个问题合并成一行(self):
        b = mk([(5, 9, 0, "张三", "合同签完了，什么时候能开始用？我们下个月就想跑起来。")])
        items = Q.extract_qa(b, {})
        self.assertEqual(len(items), 1)
        self.assertIn("什么时候能开始用", items[0].question)
        self.assertIn("想跑起来", items[0].question)

    def test_我方待办单独成行且问题列占位(self):
        b = mk([
            (5, 9, 0, "张三", "审批流能按主体拆吗"),
            (5, 9, 5, BOT, "可以，我拉一份字段映射表给你们确认"),
        ])
        items = Q.extract_qa(b, {})
        todo = [it for it in items if it.kind == "我方待办"]
        # 这条已被用作答复，不应再重复成待办
        self.assertEqual(todo, [])
        b2 = mk([(5, 9, 30, BOT, "我这边帮您配一个「海南专项」分组，今天下午能配好")])
        t2 = Q.extract_qa(b2, {})
        self.assertEqual(len(t2), 1)
        self.assertEqual(t2[0].kind, "我方待办")
        self.assertEqual(t2[0].status, QA_STATUS_PROMISED)
        self.assertEqual(t2[0].question_cell, "—")
        self.assertEqual(t2[0].answer_time, "—")

    def test_反复追问会去重(self):
        b = mk([
            (5, 9, 0, "张三", "租车标准通常情况下是多少"),
            (5, 9, 1, "张三", "租车标准通常情况下是多少?"),
        ])
        self.assertEqual(len(Q.extract_qa(b, {})), 1)

    def test_载荷消息不参与抽取也不当答复(self):
        b = mk([
            (5, 9, 0, "张三", "审批流能按主体拆吗"),
            (5, 9, 1, BOT, "10000000000000001@openim:\n<?xml version=\"1.0\"?><msg>x</msg>"),
            (5, 9, 5, BOT, "可以，按单据类型拆两条审批流"),
        ])
        items = Q.extract_qa(b, {})
        self.assertEqual(len(items), 1)
        self.assertIn("按单据类型拆两条", items[0].answer)


class TestExportIsolation(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="wxcsm_test_"))
        self.xls = self.tmp / "台账.xlsx"
        self.sums = [Summary(customer_name="测试客户", time_range="2026-08-01 ~ 2026-08-31",
                             content="占位总结")]
        self.items = mk([
            (5, 9, 0, "张三", "审批流能按主体拆吗"),
            (5, 9, 5, BOT, "可以，按单据类型拆两条审批流"),
        ])
        self.qa = Q.extract_qa(self.items, {})

    def test_两张表共存且各自追加(self):
        E.export(self.sums, self.xls, "append")
        E.export_qa(self.qa, self.xls, "append")
        wb = load_workbook(self.xls)
        self.assertEqual(wb.sheetnames, [E.SUMMARY_SHEET, E.QA_SHEET])
        self.assertEqual(wb[E.SUMMARY_SHEET].max_row - 1, 1)
        self.assertEqual(wb[E.QA_SHEET].max_row - 1, len(self.qa))
        # 再追加一轮
        E.export(self.sums, self.xls, "append")
        E.export_qa(self.qa, self.xls, "append")
        wb = load_workbook(self.xls)
        self.assertEqual(wb[E.SUMMARY_SHEET].max_row - 1, 2)
        self.assertEqual(wb[E.QA_SHEET].max_row - 1, len(self.qa) * 2)

    def test_overwrite只能重建目标表不能删掉另一张(self):
        E.export(self.sums, self.xls, "append")
        E.export_qa(self.qa, self.xls, "append")
        E.export(self.sums, self.xls, "overwrite")     # 坑 7
        wb = load_workbook(self.xls)
        self.assertIn(E.QA_SHEET, wb.sheetnames)
        self.assertEqual(wb[E.SUMMARY_SHEET].max_row - 1, 1)
        self.assertEqual(wb[E.QA_SHEET].max_row - 1, len(self.qa))

    def test_陌生工作簿原样不动(self):
        foreign = self.tmp / "别人的表.xlsx"
        wb = Workbook()
        wb.active.title = "报销明细"
        wb.active.append(["部门", "金额"])
        wb.save(foreign)
        out, _, _ = E.export(self.sums, foreign, "append")
        self.assertNotEqual(out, foreign)
        self.assertEqual(load_workbook(foreign).sheetnames, ["报销明细"])

    def test_先问答后补总结落进同一工作簿(self):
        E.export_qa(self.qa, self.xls, "append")
        out, _, _ = E.export(self.sums, self.xls, "append")
        self.assertEqual(out, self.xls)
        self.assertEqual(load_workbook(self.xls).sheetnames,
                         [E.SUMMARY_SHEET, E.QA_SHEET])


class TestOutputFormat(unittest.TestCase):
    def test_开关判定(self):
        from core import pipeline
        self.assertTrue(pipeline.wants_summary({}))
        self.assertTrue(pipeline.wants_qa({}))
        self.assertTrue(pipeline.wants_summary({"output_format": "summary"}))
        self.assertFalse(pipeline.wants_qa({"output_format": "summary"}))
        self.assertTrue(pipeline.wants_qa({"output_format": "qa"}))
        self.assertFalse(pipeline.wants_summary({"output_format": "qa"}))


if __name__ == "__main__":
    unittest.main(verbosity=2)
