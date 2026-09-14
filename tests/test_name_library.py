# -*- coding: utf-8 -*-
"""名称库单元测试。

重点不是"happy path"，而是**Excel 导入进去不该有的东西**：
员工花名册里序号/部门/手机号/表头都很容易被当成同事存进名称库。
名称库是用**子串匹配**判我方身份的（见 core/sources/wx4_live.py），
库里混进一个泛词会误伤真实客户发言，所以清洗规则值得逐条钉住。
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core import name_library as NL


class TestCleanNames(unittest.TestCase):
    def test_表头不能被当成姓名(self):
        for h in ("姓名", "名字", "员工姓名", "客户名称", "成员", "人员",
                  "手机号", "手机", "电话", "联系方式", "部门", "职位",
                  "工号", "序号", "编号", "备注", "name", "用户名"):
            self.assertEqual(NL.clean_names([h]), [], h)

    def test_数字与日期不能被当成姓名(self):
        for v in ("12", "12.0", 12, 12.5, "13800138000", "2026-09-11",
                  "2026/9/11", "100000"):
            self.assertEqual(NL.clean_names([v]), [], repr(v))

    def test_部门与组织名不能被当成姓名(self):
        for v in ("研发部", "财务科", "市场中心", "销售一组", "郑州分公司",
                  "某某科技有限公司", "项目组", "交付团队"):
            self.assertEqual(NL.clean_names([v]), [], v)

    def test_单字后缀的短名要保留(self):
        # 以"科"/"组"结尾的 2 字人名不能被当成部门名杀掉（如「张科」撞「财务科」、「李组」撞「项目组」）。
        # 名称库误杀真名的后果是"同事的答复没算成我方"，比多几条噪声更难发现。
        # 注：这里刻意用「张科」「李组」这种同形假名——换成「张三」就测不到这个后缀判断了。
        for v in ("张科", "李组"):
            self.assertEqual(NL.clean_names([v]), [v], v)

    def test_邮箱路径与长文本不能被当成姓名(self):
        for v in ("a@b.com", "C:/Users/x", "这是一段很长的备注说明文字超过二十个字符应被丢弃"):
            self.assertEqual(NL.clean_names([v]), [], v)

    def test_去空去重且保序(self):
        self.assertEqual(
            NL.clean_names(["  张三 ", "张三", "", None, "李四", "\u3000王五\u3000"]),
            ["张三", "李四", "王五"])

    def test_真实昵称形式要留下(self):
        # 真实数据里同事的备注就叫「张三 交付」，这是有效的人名形式
        self.assertEqual(NL.clean_names(["张三 交付"]), ["张三 交付"])
        self.assertEqual(NL.clean_names(["李四"]), ["李四"])


class TestMergeSplit(unittest.TestCase):
    def test_合并去重保序(self):
        self.assertEqual(NL.merge_names(["A", "B"], ["B", "C"], ["A"]),
                         ["A", "B", "C"])

    def test_合并能容忍空输入(self):
        self.assertEqual(NL.merge_names([], None, ["A"]), ["A"])

    def test_切分支持中英文标点与换行(self):
        self.assertEqual(NL.split_names("张三, 李四；王五、张科\n李组"),
                         ["张三", "李四", "王五", "张科", "李组"])

    def test_切分空串(self):
        self.assertEqual(NL.split_names(""), [])
        self.assertEqual(NL.split_names("  , ，、 "), [])


def _make_roster(path: Path) -> None:
    """造一张典型员工花名册：序号/部门/姓名/手机号，且姓名不在第一列。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "员工名单"
    ws.append(["序号", "部门", "员工姓名", "手机号", "入职日期"])
    ws.append([1, "研发部", "张三", "13800138000", "2025-01-06"])
    ws.append([2, "交付部", "李四", "13900139000", "2025-03-01"])
    ws.append([3, "交付部", "张三 交付", None, "2025-04-02"])
    ws.append([4, "研发部", "张科", "13700137000", "2025-05-03"])
    wb.save(path)


class TestReadColumns(unittest.TestCase):
    def test_从花名册里读出姓名列(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "花名册.xlsx"
            _make_roster(p)
            cols = NL.read_columns(p)

        labels = [c.label for c in cols]
        self.assertTrue(any("员工姓名" in x for x in labels), labels)
        # 姓名列在第 3 列（index 2）
        name_col = [c for c in cols if c.header == "员工姓名"][0]
        self.assertEqual(name_col.index, 2)
        self.assertEqual(name_col.names, ["张三", "李四", "张三 交付", "张科"])

    def test_序号列与部门列不会被当成人(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "花名册.xlsx"
            _make_roster(p)
            cols = {c.header: c for c in NL.read_columns(p)}

        # 序号列全是数字 → 清洗后空
        self.assertEqual(cols["序号"].names, [])
        # 部门列有重复且都是部门 → 清洗后空
        self.assertEqual(cols["部门"].names, [])
        # 手机号列 → 空
        self.assertEqual(cols["手机号"].names, [])

    def test_能猜中姓名列(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "花名册.xlsx"
            _make_roster(p)
            cols = NL.read_columns(p)

        guess = NL.guess_column(cols)
        self.assertIsNotNone(guess)
        self.assertEqual(guess.header, "员工姓名")

    def test_没有表头时靠名字数量猜(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "无表头.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["张三"])
            ws.append(["李四"])
            ws.append(["王五"])
            wb.save(p)
            guess = NL.guess_column(NL.read_columns(p))

        self.assertIsNotNone(guess)
        self.assertEqual(guess.names, ["张三", "李四", "王五"])

    def test_空文件不报错(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "空.xlsx"
            Workbook().save(p)
            cols = NL.read_columns(p)

        self.assertEqual([c for c in cols if c.names], [])
        self.assertIsNone(NL.guess_column(cols))

    def test_多sheet都能读到(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "两个表.xlsx"
            wb = Workbook()
            wb.active.title = "A表"
            wb.active.append(["姓名"])
            wb.active.append(["张三"])
            ws2 = wb.create_sheet("B表")
            ws2.append(["姓名"])
            ws2.append(["李四"])
            wb.save(p)
            cols = NL.read_columns(p)

        sheets = {c.sheet for c in cols}
        self.assertEqual(sheets, {"A表", "B表"})


class TestConfigKey(unittest.TestCase):
    def test_名称库是登记的配置键(self):
        # config.save() 只写 DEFAULTS 里有的键，新键忘了登记就会"存了没生效"
        from core import config

        self.assertIn("self_name_library", config.DEFAULTS)
        self.assertEqual(config.DEFAULTS["self_name_library"], [])


class TestLooksLikePerson(unittest.TestCase):
    def test_可疑词能被识别(self):
        for v in ("研发部", "13800138000", "员工姓名", "2026-09-11", ""):
            self.assertFalse(NL.looks_like_person(v), v)
            self.assertEqual(NL.suspect_names([v]), [v], v)

    def test_正常姓名通过(self):
        for v in ("张三", "李四", "张三 交付", "张科", "王五"):
            self.assertTrue(NL.looks_like_person(v), v)
            self.assertEqual(NL.suspect_names([v]), [], v)

    def test_手动输入不被过滤只被提示(self):
        # 这是刻意的设计：手动敲的名字应当尊重，静默丢弃会让人以为"加了没生效"
        self.assertEqual(NL.split_names("张三,研发部"), ["张三", "研发部"])
        self.assertEqual(NL.suspect_names(NL.split_names("张三,研发部")), ["研发部"])


class TestConfigCommand(unittest.TestCase):
    """config 子命令：名称库要能从命令行存取（脚本化/WorkBuddy 编排要用）。"""

    def setUp(self):
        import tempfile

        from core import config as cfgmod

        self._tmp = tempfile.TemporaryDirectory()
        self._old_path = cfgmod.CONFIG_PATH
        cfgmod.CONFIG_PATH = Path(self._tmp.name) / "config.json"

    def tearDown(self):
        from core import config as cfgmod

        cfgmod.CONFIG_PATH = self._old_path
        self._tmp.cleanup()

    def _args(self, **kw):
        class A:
            show = False
            add_library = None
            remove_library = None
            self_names = None
            import_library = None
            export_template = None
            force = False
            sheet = None
            column = None
            json = False

        a = A()
        for k, v in kw.items():
            setattr(a, k, v)
        return a

    def _cfg(self):
        from core import config as cfgmod

        return cfgmod.load()

    def test_增删名称库并保留选用名单(self):
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS, "self_names": ["张三"]})
        self.assertEqual(cli.cmd_config(self._args(add_library="张三,李四")), 0)
        self.assertEqual(self._cfg()["self_name_library"], ["张三", "李四"])
        self.assertEqual(self._cfg()["self_names"], ["张三"], "不该动选用名单")

        self.assertEqual(cli.cmd_config(self._args(remove_library="李四")), 0)
        self.assertEqual(self._cfg()["self_name_library"], ["张三"])

    def test_删除也清理选用名单(self):
        # 否则会出现"库里没了但还在用"的矛盾状态
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS, "self_name_library": ["张三", "李四"],
                     "self_names": ["张三", "李四"]})
        cli.cmd_config(self._args(remove_library="张三"))
        self.assertEqual(self._cfg()["self_name_library"], ["李四"])
        self.assertEqual(self._cfg()["self_names"], ["李四"])

    def test_从excel导入进名称库(self):
        import cli

        from core import config as cfgmod

        cfgmod.save(dict(cfgmod.DEFAULTS))
        p = Path(self._tmp.name) / "名册.xlsx"
        _make_roster(p)
        self.assertEqual(cli.cmd_config(self._args(import_library=str(p))), 0)
        lib = self._cfg()["self_name_library"]
        for n in ("张三", "李四", "张三 交付", "张科"):
            self.assertIn(n, lib)
        for bad in ("研发部", "交付部", "13800138000"):
            self.assertNotIn(bad, lib)

    def test_导入以文件为准会移除库里多出来的人(self):
        """命令行要与界面一致：文件里没有的人会从库里移除。

        会删人就要求显式 --force（界面上对应那个确认框）：命令行跑批没人看着，
        默认直接拒绝，而不是静默把库清掉。
        """
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS,
                     "self_name_library": ["王五", "张三"],
                     "self_names": ["王五", "张三"]})
        p = Path(self._tmp.name) / "名册.xlsx"
        _make_roster(p)

        self.assertEqual(cli.cmd_config(self._args(import_library=str(p))), 1,
                         "会删人却没给 --force，应当拒绝执行")
        self.assertEqual(self._cfg()["self_name_library"], ["王五", "张三"],
                         "拒绝执行时不该动名称库")
        self.assertEqual(self._cfg()["self_names"], ["王五", "张三"],
                         "拒绝执行时不该动选用名单")

        self.assertEqual(
            cli.cmd_config(self._args(import_library=str(p), force=True)), 0)
        lib = self._cfg()["self_name_library"]
        self.assertNotIn("王五", lib, "文件里没有的人应当被移除")
        for n in ("张三", "李四", "张三 交付", "张科"):
            self.assertIn(n, lib, f"{n} 应当进库")
        # 库里没了却还留在选用名单里是矛盾状态（同 --remove-library）
        self.assertNotIn("王五", self._cfg()["self_names"])
        self.assertIn("张三", self._cfg()["self_names"])

    def test_导入纯新增不需要确认(self):
        """只增不删没有数据损失，不该拦——否则脚本每次都得写 --force。"""
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS, "self_name_library": ["张三"]})
        p = Path(self._tmp.name) / "名册.xlsx"
        _make_roster(p)
        self.assertEqual(cli.cmd_config(self._args(import_library=str(p))), 0)

    def test_add和import不能同时用(self):
        """导入是"以文件为准"的替换，会冲掉刚 add 的人，所以直接拒绝。"""
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS, "self_name_library": ["张三"]})
        p = Path(self._tmp.name) / "名册.xlsx"
        _make_roster(p)
        code = cli.cmd_config(self._args(add_library="王五", import_library=str(p),
                                         force=True))
        self.assertEqual(code, 2, "矛盾的组合应当被拒")
        self.assertEqual(self._cfg()["self_name_library"], ["张三"], "不该有副作用")

    def test_导出模版带出名称库(self):
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS, "self_name_library": ["张三", "李四"]})
        out = Path(self._tmp.name) / "模版.xlsx"
        self.assertEqual(cli.cmd_config(self._args(export_template=str(out))), 0)
        self.assertTrue(out.exists(), "模版没写出来")
        names = [n for c in NL.read_columns(out) for n in c.names]
        self.assertEqual(names, ["张三", "李四"], "模版应带出名称库的人")
        self.assertNotIn("姓名", names, "表头不能被当成人名带出去")

    def test_导入加导出拿到的是更新后的库(self):
        """导出排在最后执行：导出的必须是"刚改完的库"，不是改之前的。"""
        import cli

        from core import config as cfgmod

        cfgmod.save({**cfgmod.DEFAULTS, "self_name_library": ["王五"]})
        p = Path(self._tmp.name) / "名册.xlsx"
        _make_roster(p)
        out = Path(self._tmp.name) / "模版.xlsx"
        self.assertEqual(
            cli.cmd_config(self._args(import_library=str(p), force=True,
                                      export_template=str(out))), 0)
        names = [n for c in NL.read_columns(out) for n in c.names]
        self.assertNotIn("王五", names, "导出的是更新后的库，王五已被移除")
        self.assertIn("张科", names)

    def test_指定列号导入(self):
        import cli

        from core import config as cfgmod

        cfgmod.save(dict(cfgmod.DEFAULTS))
        p = Path(self._tmp.name) / "名册.xlsx"
        _make_roster(p)
        # 第 3 列是姓名列
        cli.cmd_config(self._args(import_library=str(p), column=3))
        self.assertEqual(self._cfg()["self_name_library"],
                         ["张三", "李四", "张三 交付", "张科"])

    def test_没有姓名列时报错而不是静默成功(self):
        import cli

        from core import config as cfgmod
        from openpyxl import Workbook

        cfgmod.save(dict(cfgmod.DEFAULTS))
        p = Path(self._tmp.name) / "只有数字.xlsx"
        wb = Workbook()
        wb.active.append(["序号"])
        wb.active.append([1])
        wb.active.append([2])
        wb.save(p)
        self.assertEqual(cli.cmd_config(self._args(import_library=str(p))), 1)
        self.assertEqual(self._cfg()["self_name_library"], [])


class TestHeaderDecoration(unittest.TestCase):
    """真实模板的表头常带装饰（姓名*、姓名（必填））。

    不处理的话 _HEADER_WORD 匹配不上——因为要求整串等于或以「姓名」结尾——
    结果**表头本身会被当成一个人名导进名称库**。
    """

    def test_剥掉装饰尾巴(self):
        for raw, want in (("姓名（必填）", "姓名"), ("姓名*", "姓名"), ("姓名 *", "姓名"),
                          ("员工姓名（必填）", "员工姓名"), ("姓名[必填]", "姓名"),
                          ("姓名", "姓名"), ("员工姓名", "员工姓名")):
            self.assertEqual(NL.undecorate_header(raw), want, raw)

    def test_带装饰的表头不会被当成姓名(self):
        for v in ("姓名（必填）", "姓名*", "员工姓名（必填）", "姓名[必填]",
                  "姓名（必填项）", "成员（必填）"):
            self.assertEqual(NL.clean_names([v]), [], v)

    def test_带括号的真名不受影响(self):
        # 「李四（销售）」剥掉装饰后是「李四」，比对不匹配表头词，应当保留
        self.assertEqual(NL.clean_names(["李四（销售）"]), ["李四（销售）"])

    def test_装饰表头的列仍能被猜中(self):
        col = NL.NameColumn(sheet="S", index=0, header="姓名（必填）", names=[])
        self.assertTrue(col.likely_name_column)


class TestTemplate(unittest.TestCase):
    """导出导入模版：必填项标红，且空模版导进来不能产生任何名字。"""

    def setUp(self):
        import tempfile

        self._tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self._tmp.cleanup()

    def _make(self):
        return NL.write_template(Path(self._tmp.name) / "模版.xlsx")

    def test_模版表头与必填标红(self):
        from openpyxl import load_workbook

        p = self._make()
        self.assertTrue(p.exists())
        wb = load_workbook(p)
        ws = wb.active
        self.assertEqual(ws.title, NL.TEMPLATE_SHEET)
        heads = [c.value for c in ws[1]]
        self.assertEqual(heads, ["姓名（必填）", "部门（可选）", "手机号（可选）", "备注（可选）"])

        def rgb(cell):
            col = cell.font.color
            return getattr(col, "rgb", None) if col is not None else None

        self.assertTrue(str(rgb(ws["A1"])).endswith(NL.TEMPLATE_REQUIRED_COLOR[-6:]),
                        f"姓名表头没标红：{rgb(ws['A1'])}")
        self.assertTrue(ws["A1"].font.bold, "必填表头该是粗体")
        for ref in ("B1", "C1", "D1"):
            self.assertFalse(str(rgb(ws[ref])).endswith(NL.TEMPLATE_REQUIRED_COLOR[-6:]),
                             f"{ref} 是可选项，不该标红")
        wb.close()

    def test_空模版导入不产生任何名字(self):
        # 最要紧的一条：表头「姓名（必填）」若被当成真人名存进名称库，
        # 而名称库是按子串匹配判我方身份的，可能误伤真实客户发言。
        cols = NL.read_columns(self._make())
        names = [n for c in cols for n in c.names]
        self.assertEqual(names, [], f"空模版导出了假名字：{names}")

    def test_模版只有表头没有示例数据(self):
        # 不传名称库时（相当于要一份空表）只有表头，刻意不造示例行——
        # 用户会直接填在示例下面，示例名字就被一起导进去了。
        # 传了名称库时写的是**用户自己的**库成员：他们本来就在库里，不是伪数据。
        from openpyxl import load_workbook

        wb = load_workbook(self._make())
        ws = wb.active
        self.assertEqual(ws.max_row, 1, "模版不该带示例数据行")
        wb.close()

    def test_模版能把名称库的人带出来(self):
        # 导出→在 Excel 里改→导入，是个闭环：模版必须带出当前库里的人，
        # 否则用户只能对着空表回忆库里都有谁。
        p = NL.write_template(Path(self._tmp.name) / "带库.xlsx",
                              ["张三", "李四", "王五"])
        from openpyxl import load_workbook

        wb = load_workbook(p)
        ws = wb.active
        self.assertEqual([ws.cell(row=r, column=1).value for r in range(1, 5)],
                         ["姓名（必填）", "张三", "李四", "王五"])
        wb.close()
        # 表头不能被当成一个人带进去
        got = NL.guess_column(NL.read_columns(p))
        self.assertIsNotNone(got)
        self.assertEqual(got.names, ["张三", "李四", "王五"])

    def test_带名称库时不过滤手动加的名字(self):
        # 手动添加的名字按设计就不过滤（"你明确敲进来的名字应当尊重"）。
        # 导出时若再 filter 一遍，「研发部」这种用户手敲的名字会从模版里消失，
        # 他改完导入回来就发现人没了——等于偷偷删数据。
        p = NL.write_template(Path(self._tmp.name) / "不过滤.xlsx", ["研发部", "张三"])
        from openpyxl import load_workbook

        wb = load_workbook(p)
        ws = wb.active
        self.assertEqual([ws.cell(row=r, column=1).value for r in range(1, 4)],
                         ["姓名（必填）", "研发部", "张三"])
        wb.close()
        # 反向确认这个名字确实会被导入侧的 clean_names 剔掉（所以"不过滤"是必需的）
        self.assertEqual(NL.clean_names(["研发部"]), [])

    def test_填好模版后能正常导入(self):
        from openpyxl import load_workbook

        p = self._make()
        wb = load_workbook(p)
        ws = wb.active
        for i, n in enumerate(("张三", "李四", "王五"), start=2):
            ws.cell(row=i, column=1, value=n)
        wb.save(p)
        wb.close()

        cols = NL.read_columns(p)
        got = NL.guess_column(cols)
        self.assertIsNotNone(got, "填好后应当猜得出姓名列")
        self.assertEqual(got.names, ["张三", "李四", "王五"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
