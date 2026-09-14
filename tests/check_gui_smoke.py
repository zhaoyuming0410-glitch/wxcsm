# -*- coding: utf-8 -*-
"""GUI 冒烟测试：不开 mainloop，直接把窗口建起来、驱动关键路径、检查状态。

覆盖第③步新增的东西——这些是"只有点开界面才会发现"的错误：
  1. 窗口能否建起来（grid 选项写错、引用不存在的控件都会在这里炸）
  2. 输出形式切换时结果区列是否正确切换、字数/阶段是否按预期置灰
  3. 问答条目能否正确渲染（含状态标色）
  4. 双击/修订能否打开对应编辑器（总结 / 问答各自一个）
  5. 导出走 pipeline.export_results 后，两张 sheet 是否都落到文件里

注意：_sync_cfg 会写真实配置，所以全程备份并还原 config.json。
"""
from __future__ import annotations

import io
import sys
import traceback
from datetime import datetime
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import tkinter as tk

from core.models import (QA_STATUS_DONE, QA_STATUS_PENDING, QA_STATUS_PROMISED,
                         Contact, QaItem, Summary)
from core.self_discover import Candidate

import app as appmod

CFG = Path.home() / ".wxcsm" / "config.json"
BACKUP = CFG.read_text(encoding="utf-8") if CFG.exists() else None

FAILS: list[str] = []


class FakeMB:
    """替掉 messagebox：既避免模态框阻塞测试，又能断言弹了什么。"""

    def __init__(self) -> None:
        self.calls: list[tuple] = []
        # 导入名称库前的"确认更新"弹窗：默认点确认，单个用例可以改成 False
        # 来验证"取消后名称库原样不动"。
        self.okcancel_answer = True

    def showinfo(self, *a, **k):
        self.calls.append(("info", a))

    def showwarning(self, *a, **k):
        self.calls.append(("warn", a))

    def showerror(self, *a, **k):
        self.calls.append(("error", a))

    def askyesno(self, *a, **k):
        self.calls.append(("ask", a))
        return False

    def askokcancel(self, *a, **k):
        self.calls.append(("okcancel", a))
        return self.okcancel_answer

    def last(self, kind: str):
        """最近一次某类弹窗的参数，没有则 None。"""
        hits = [a for k, a in self.calls if k == kind]
        return hits[-1] if hits else None


def check(name: str, fn) -> None:
    try:
        fn()
        print(f"  [OK]    {name}")
    except Exception as e:                                          # noqa: BLE001
        FAILS.append(name)
        print(f"  [FAIL]  {name}: {type(e).__name__}: {e}")
        traceback.print_exc(limit=4)


def qa(kind, question, answer, status, ask, ans=None, name="某客户"):
    return QaItem(customer_name=name, kind=kind, question=question, answer=answer,
                  status=status, raw=f"【原话】{question}", ask_ts=ask, answer_ts=ans,
                  order=0)


def main() -> int:
    mb = FakeMB()
    appmod.messagebox = mb                                          # type: ignore
    a = appmod.App()
    a.update()

    print("【1】窗口构建")
    check("主窗口创建成功", lambda: a.winfo_exists())
    check("④ 有「生成内容」单选 var_format", lambda: a.var_format.get() in
          ("both", "summary", "qa"))
    check("④ 有「我方成员」输入框", lambda: isinstance(a.var_self_names.get(), str))
    check("④ 有「选择…」按钮", lambda: a.btn_self_pick.cget("text") == "选择…")
    check("结果区按钮已改名为「修订选中条目」逻辑",
          lambda: a.btn_edit.cget("text") in ("修订选中总结", "修订选中条目"))
    check("结果区有横向滚动条", lambda: a.sb_x.winfo_exists())

    print("【2】输出形式切换")
    a.var_format.set("qa")
    a._update_format_ui()
    a.update()
    check("选「仅问答清单」时字数框置灰",
          lambda: str(a.ent_min.cget("state")) == "disabled")
    check("选「仅问答清单」时客户阶段置灰",
          lambda: str(a.cbo_stage.cget("state")) == "disabled")
    check("字数说明带「仅沟通总结用」提示",
          lambda: "仅沟通总结用" in a.lbl_chars_end.cget("text"))
    a.var_format.set("summary")
    a._update_format_ui()
    a.update()
    check("切回「仅总结」时字数框恢复可用",
          lambda: str(a.ent_min.cget("state")) == "normal")
    check("客户阶段恢复 readonly",
          lambda: str(a.cbo_stage.cget("state")) == "readonly")

    print("【3】问答清单渲染")
    a.summaries = [Summary(customer_name="某客户", time_range="2026-08-01 ~ 2026-08-31",
                           content="总结正文", msg_count=12, engine="offline")]
    a.qa_items = [
        qa("提问", "能不能改配置？", "已经改好了。", QA_STATUS_DONE,
           datetime(2026, 8, 1, 10, 0), datetime(2026, 8, 1, 11, 0)),
        qa("报错", "导不出来", "", QA_STATUS_PENDING, datetime(2026, 8, 2, 9, 0)),
        qa("我方待办", "", "明天给他配置表", QA_STATUS_PROMISED,
           datetime(2026, 8, 3, 9, 0)),
    ]
    a.var_format.set("both")
    a._update_format_ui()
    a.update()
    check("两者都要且有条目 → 显示问答形态",
          lambda: a._res_shape == "qa", )
    check("问答列数 = 8",
          lambda: len(a.tv_res.cget("columns")) == 8)
    check("问答行数 = 3", lambda: len(a.tv_res.get_children()) == 3)
    check("表头是白话列名",
          lambda: a.tv_res.heading("question")["text"] == "客户问题")
    check("未答复行标红 tag",
          lambda: "pend" in a.tv_res.item("1", "tags"))
    check("我方待办行标橙 tag",
          lambda: "promise" in a.tv_res.item("2", "tags"))
    check("我方待办行问题列显示占位符 —",
          lambda: a.tv_res.item("2", "values")[3] == "—")
    check("未答复行答复时间显示 —",
          lambda: a.tv_res.item("1", "values")[5] == "—")
    check("底部统计含未答复/待办数",
          lambda: "未答复 1" in a.lbl_res.cget("text")
                  and "我方待办 1" in a.lbl_res.cget("text"))

    print("【4】形态切换回总结")
    a.var_format.set("summary")
    a._update_format_ui()
    a.update()
    check("切到「仅总结」→ 显示总结形态", lambda: a._res_shape == "summary")
    check("总结列数 = 4", lambda: len(a.tv_res.cget("columns")) == 4)
    check("总结行数 = 1", lambda: len(a.tv_res.get_children()) == 1)
    check("按钮文案跟着变",
          lambda: a.btn_edit.cget("text") == "修订选中总结")

    print("【5】行编辑器")
    a.var_format.set("qa")
    a._update_format_ui()
    a.update()
    a.tv_res.selection_set("0")

    def open_qa_editor():
        a._edit_qa()
        a.update()
        tops = [w for w in a.winfo_children() if isinstance(w, tk.Toplevel)]
        assert tops, "没有弹出编辑器窗口"
        t = tops[-1]
        title = t.title()
        t.destroy()
        a.update()
        assert "修订问答条目" in title, title

    check("双击问答行能打开问答编辑器", open_qa_editor)

    def open_summary_editor():
        a.var_format.set("summary")
        a._update_format_ui()
        a.update()
        a.tv_res.selection_set("0")
        a._edit_row()
        a.update()
        tops = [w for w in a.winfo_children() if isinstance(w, tk.Toplevel)]
        assert tops, "没有弹出编辑器窗口"
        t = tops[-1]
        title = t.title()
        t.destroy()
        a.update()
        assert "修订总结" in title, title

    check("双击总结行能打开总结编辑器", open_summary_editor)

    print("【6】我方成员选择窗口")
    a.contacts = [Contact(cid="c1", name="某交付群", kind="group", msg_count=10),
                  Contact(cid="c2", name="某客户", kind="friend", msg_count=5)]
    a.picked = {"c1", "c2"}
    a.var_self_names.set("张三, 李四")
    a.self_candidates = []

    def open_picker():
        a._open_self_picker()
        a.update()
        tops = [w for w in a.winfo_children() if isinstance(w, tk.Toplevel)]
        assert tops, "没有弹出我方成员窗口"
        t = tops[-1]
        # 找到候选表，核对预填写：两名已配置成员应在列表里且已勾选
        trees = [w for w in t.winfo_children() if isinstance(w, tk.Frame)]
        found = None
        for f in t.winfo_children():
            for sub in getattr(f, "winfo_children", lambda: [])():
                if isinstance(sub, tk.ttk.Treeview):
                    found = sub
        assert found is not None, "没有找到候选表"
        rows = [found.item(i, "values") for i in found.get_children()]
        names = [r[1] for r in rows]
        assert "张三" in names and "李四" in names, f"预填写名单缺少已配置成员：{names}"
        # 只断言「已配置的那两位」是勾上的。名称库里可能还有别的同事（长期存着、
        # 这次不一定要用），它们本来就该是不勾的——写成 all(...) 会误判。
        d = {r[1]: r for r in rows}
        for who in ("张三", "李四"):
            assert d[who][0] == "☑", f"已配置成员应默认勾选：{d[who]}"
        t.destroy()
        a.update()

    check("选择我方成员窗口：预填写已配置成员且默认勾选", open_picker)

    # ---- 名称库（第③步追加）：库里存人、手动添加、Excel 导入、确定入库 ----
    def find_widgets(root, cls, text=None):
        out = []
        stack = [root]
        while stack:
            w = stack.pop()
            if isinstance(w, cls) and (text is None or
                                       str(w.cget("text") if "text" in w.keys() else "") == text):
                out.append(w)
            stack.extend(w.winfo_children())
        return out

    def picker_window():
        tops = [w for w in a.winfo_children() if isinstance(w, tk.Toplevel)]
        assert tops, "没有弹出我方成员窗口"
        return tops[-1]

    def rows_of(t):
        tv = find_widgets(t, tk.ttk.Treeview)
        assert tv, "没有找到候选表"
        return tv[0], [tv[0].item(i, "values") for i in tv[0].get_children()]

    def click_row(tree, name_or_iid):
        """真的去点某一行：按行坐标发 Button-1。

        不能只 selection_set 再发一个空事件——空事件的 x/y 都是 0（落在表头上），
        按坐标取行的实现会正确地忽略它，测试就成了空转。
        """
        iid = name_or_iid
        for i in tree.get_children():
            if i == name_or_iid or tree.item(i, "values")[1] == name_or_iid:
                iid = i
                break
        bb = tree.bbox(iid)
        assert bb, f"行 {iid} 不在可见区域，点不到（bbox 为空）"
        x, y, w, h = bb
        tree.event_generate("<Button-1>", x=x + w // 2, y=y + h // 2, when="now")

    def sum_text(t):
        """读「已选 N 人…」汇总行。它挂的是 textvariable，cget("text") 是空的。"""
        out = []
        for w in find_widgets(t, tk.ttk.Label):
            var = str(w.cget("textvariable") or "")
            if not var:
                continue
            try:
                txt = str(t.getvar(var))
            except tk.TclError:
                continue
            if "已选" in txt:
                out.append(txt)
        return out[-1] if out else ""

    a.cfg["self_name_library"] = ["张三", "李四"]
    a.var_self_names.set("王五")
    a.self_candidates = []

    a._open_self_picker()
    a.update()
    t = picker_window()
    tv, rows = rows_of(t)
    by_name = {r[1]: r for r in rows}

    def lib_rows_checked():
        assert "张三" in by_name and "李四" in by_name, f"名称库成员没进列表：{list(by_name)}"
        assert "王五" in by_name, "已配置成员没进列表"
        # 库里的人默认**不勾**（这次不一定用），已配置的那位才是勾上的
        assert by_name["王五"][0] == "☑", f"已配置成员应勾选：{by_name['王五']}"
        assert by_name["张三"][0] == "☐", f"库成员默认不该勾：{by_name['张三']}"
        assert "名称库" in by_name["张三"][4], f"来源列没标名称库：{by_name['张三']}"

    check("名称库：库成员进列表但不默认勾选，来源标注正确", lib_rows_checked)

    def manual_add_enters_library():
        ent = find_widgets(t, tk.ttk.Entry)
        assert ent, "没有找到手动输入框"
        ent[0].insert(0, "张科")
        # 必须给焦点 + when="now"，否则 <Return> 不会被派发到绑定函数上
        ent[0].focus_force()
        ent[0].event_generate("<Return>", when="now")
        a.update()
        _, rows2 = rows_of(t)
        d = {r[1]: r for r in rows2}
        assert "张科" in d, f"手动添加的人没进列表：{list(d)}"
        assert d["张科"][0] == "☑", f"手动添加应默认勾选：{d['张科']}"
        assert "手动添加" in d["张科"][4], f"来源没标手动添加：{d['张科']}"

    check("名称库：手动输入添加并默认勾选", manual_add_enters_library)

    # --- Excel 导入：真的走 read_columns + 选列对话框 + 导入按钮 ---
    import tempfile
    from openpyxl import Workbook

    roster = Path(tempfile.mkdtemp(prefix="wxcsm_names_")) / "花名册.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "部门", "员工姓名", "手机号"])
    ws.append([1, "研发部", "李组", "13800138000"])
    ws.append([2, "交付部", "李工", "13900139000"])
    wb.save(roster)

    real_askopen = appmod.filedialog.askopenfilename
    appmod.filedialog.askopenfilename = lambda *x, **k: str(roster)

    def excel_import():
        btn = find_widgets(t, tk.ttk.Button, "导入 Excel…")
        assert btn, "没有找到「导入 Excel…」按钮"
        before = {r[1] for r in rows_of(t)[1]}
        assert {"张三", "李四", "王五", "张科"} <= before, f"导入前的名称库不对：{before}"
        btn[0].invoke()
        a.update()
        # 应弹出「选择姓名列」对话框，且默认猜中「员工姓名」列。
        # 注意：这个对话框的 parent 是**选择窗口** t，不是主窗口 a。
        tops = [w for w in t.winfo_children() if isinstance(w, tk.Toplevel)]
        assert tops, "没有弹出选择姓名列的对话框"
        dlg = tops[-1]
        combo = find_widgets(dlg, tk.ttk.Combobox)
        assert combo, "没有找到列选择下拉框"
        assert "员工姓名" in combo[0].get(), f"没猜中姓名列：{combo[0].get()}"
        go = find_widgets(dlg, tk.ttk.Button, "导入这些名字")
        assert go, "没有找到导入按钮"
        go[0].invoke()
        a.update()
        # 导入以文件为准 → 会删人，所以必须先把增删名单摆出来让人确认
        asked = mb.last("okcancel")
        assert asked, "会删人却没弹确认框"
        text = "\n".join(str(x) for x in asked)
        for who in ("张三", "李四", "王五", "张科"):
            assert who in text, f"确认框没列出将被移除的人 {who}：{text}"
        assert "李组" in text and "李工" in text, f"确认框没列出将新增的人：{text}"
        _, rows3 = rows_of(t)
        d = {r[1]: r for r in rows3}
        for n in ("李组", "李工"):
            assert n in d, f"{n} 没被导入：{list(d)}"
            assert d[n][0] == "☑", f"导入的人应默认勾选：{d[n]}"
        # 最要紧的一条：以文件为准 = 库里不在文件里的人被移除（不是"只增不删"）
        for gone in ("张三", "李四", "王五", "张科"):
            assert gone not in d, f"{gone} 不在导入的文件里，应当从名称库移除：{list(d)}"
        # 序号列/部门列/手机号列不该有人混进来
        for bad in ("研发部", "交付部", "13800138000", "1", "2"):
            assert bad not in d, f"非姓名列的内容被导入了：{bad}"

    try:
        check("名称库：导入以文件为准（移除库里多出来的人，且删人前先确认）", excel_import)
    finally:
        appmod.filedialog.askopenfilename = real_askopen

    def excel_import_cancel_keeps_library():
        """确认框点取消 → 名称库原样不动。

        导错文件（比如翻出一份旧模版）会把库里的人清掉，名称库是长期资产，
        误清一次就得重新导一遍——所以取消必须是真的什么都没发生。
        """
        appmod.filedialog.askopenfilename = lambda *x, **k: str(roster)
        try:
            a.self_candidates = []
            a._open_self_picker()
            a.update()
            t6 = picker_window()
            # 钉住候选表本身：列选择对话框里也有一个 Treeview（单列），取消后
            # 那个对话框还开着，再用 rows_of(t6) 会先找到它。
            tv6 = rows_of(t6)[0]
            before = {tv6.item(i, "values")[1] for i in tv6.get_children()}
            find_widgets(t6, tk.ttk.Button, "导入 Excel…")[0].invoke()
            a.update()
            dlg6 = [w for w in t6.winfo_children() if isinstance(w, tk.Toplevel)][-1]
            mb.okcancel_answer = False
            find_widgets(dlg6, tk.ttk.Button, "导入这些名字")[0].invoke()
            a.update()
            after = {tv6.item(i, "values")[1] for i in tv6.get_children()}
            assert after == before, f"点取消后名称库被改了：{before} → {after}"
            t6.destroy()          # 连同里面那个列选择对话框一起关掉
            a.update()
        finally:
            mb.okcancel_answer = True
            appmod.filedialog.askopenfilename = real_askopen

    check("名称库：导入确认框点取消 → 名称库原样不动",
          excel_import_cancel_keeps_library)

    def excel_import_without_use():
        """取消「导入后同时勾选为本次使用」→ 只入库，不设为本次使用。

        这条是为大批量花名册准备的：整个部门/公司导进来时，全勾上会让所有人
        都变成「我方成员」，而「我方」认错会把客户发言当成我方答复。
        """
        appmod.filedialog.askopenfilename = lambda *x, **k: str(roster)
        try:
            a.self_candidates = []
            a._open_self_picker()
            a.update()
            t5 = picker_window()
            find_widgets(t5, tk.ttk.Button, "导入 Excel…")[0].invoke()
            a.update()
            dlg5 = [w for w in t5.winfo_children() if isinstance(w, tk.Toplevel)][-1]

            # 对话框装得下新加的那行提示吗
            dlg5.update_idletasks()
            need_h = dlg5.winfo_reqheight()
            got_h = dlg5.winfo_height()
            assert need_h <= got_h, f"导入对话框装不下内容：需要 {need_h}px，实际 {got_h}px"

            boxes = find_widgets(dlg5, tk.ttk.Checkbutton)
            assert boxes, "没有找到「导入后同时勾选为本次使用」勾选项"
            boxes[0].invoke()                      # 默认是勾上的，点一下取消
            a.update()
            find_widgets(dlg5, tk.ttk.Button, "导入这些名字")[0].invoke()
            a.update()

            # 关掉窗口让「确定」把它写进名称库
            find_widgets(t5, tk.ttk.Button, "确定")[0].invoke()
            a.update()
            lib = a.cfg.get("self_name_library") or []
            assert "李组" in lib, f"取消勾选后仍应入库：{lib}"
            sel = [x.strip() for x in a.var_self_names.get().split(",") if x.strip()]
            assert "李组" not in sel, f"取消勾选后不该成为本次使用：{sel}"
        finally:
            appmod.filedialog.askopenfilename = real_askopen

    check("名称库：取消「同时勾选」→ 只入库、不设为本次使用", excel_import_without_use)

    def confirm_persists_library():
        # t 窗口刚被"以文件为准"更新过：库里只剩导入文件里的人（李组/李工）且都勾着。
        # 点确定 → 勾选生效、库里就是这些人，且被移除的人不会复活。
        tv.set("0", "chk", tv.item("0", "values")[0])   # 占位，避免空选中
        names_before = list(a.cfg["self_name_library"])
        okbtn = find_widgets(t, tk.ttk.Button, "确定")
        assert okbtn, "没有找到确定按钮"
        okbtn[0].invoke()
        a.update()
        sel = [x.strip() for x in a.var_self_names.get().split(",") if x.strip()]
        assert "李组" in sel and "李工" in sel, f"勾选的成员没写回去：{sel}"
        lib = a.cfg.get("self_name_library") or []
        assert sorted(lib) == ["李工", "李组"], f"名称库应以导入的文件为准：{lib}"
        # 被导入移除的人不能因为"刚才还勾着"就复活：点确定时"勾选的人顺手入库"
        # （merge_names），所以移除时必须连勾选一起清掉。
        for gone in ("张三", "李四", "王五", "张科"):
            assert gone not in lib, f"被移除的 {gone} 又回到名称库了：{lib}"
            assert gone not in sel, f"被移除的 {gone} 仍在本次选用里：{sel}"
        assert set(names_before) <= set(lib), "库里原有成员不该丢失"

    check("名称库：确定后勾选生效、名称库保持「以文件为准」、被移除的不复活",
          confirm_persists_library)

    def click_toggles_clicked_row():
        """点哪一行就勾哪一行。

        老实现读的是 tv.selection()，而选中项要等 Treeview 的**类绑定**才更新，
        控件自身的 <Button-1> 先执行——于是读到的是"上一次点的行"：
        换一行点会勾错行，每行还得点两次。这里特意先选中另一行再点这一行，
        就是为了复现那个错位。
        """
        a.self_candidates = []
        a._open_self_picker()
        a.update()
        tw = picker_window()
        tvw = rows_of(tw)[0]
        try:
            ids = list(tvw.get_children())
            assert len(ids) >= 2, f"表里至少要有两行才能验证：{len(ids)}"
            first, second = ids[0], ids[1]
            before = {i: tvw.item(i, "values")[0] for i in ids}

            tvw.selection_set(first)        # 制造"选中项 ≠ 要点的行"
            a.update()
            click_row(tvw, second)
            a.update()

            after = {i: tvw.item(i, "values")[0] for i in ids}
            changed = [i for i in ids if before[i] != after[i]]
            assert changed == [second], (
                f"点第 2 行却改动了 {changed}（应只有 {[second]}）——"
                "又变成了「读到上一次选中行」的老毛病")
        finally:
            tw.destroy()
            a.update()

    check("勾选：点哪行就勾哪行（不会勾成上一次点的那行）", click_toggles_clicked_row)

    def select_all_and_none():
        """全选 / 全不选，以及汇总行要跟着变。"""
        # 放一个"扫描候选"进去：它不在名称库里，勾上它就该出安全提醒
        a.self_candidates = [Candidate(name="可疑客户", cid="wxid_kehu",
                                       msg_count=9, reasons=["扫描候选"])]
        try:
            a._open_self_picker()
            a.update()
            tw = picker_window()
            tvw = rows_of(tw)[0]
            try:
                ids = list(tvw.get_children())
                assert ids, "表是空的"
                # 多出来的「批量」行不能把候选表挤到不可用
                h = tvw.winfo_height()
                assert h >= 120, f"「批量」行把候选表挤得太小了：只有 {h}px"
                # 先弄成"部分勾选"，免得本来就已经全勾、全选看不出效果
                click_row(tvw, ids[0])
                a.update()
                marks = [tvw.item(i, "values")[0] for i in ids]
                assert "☑" in marks and "☐" in marks, \
                    f"需要「部分勾选」的前置状态，实际是 {marks}"

                find_widgets(tw, tk.ttk.Button, "全选")[0].invoke()
                a.update()
                got = [tvw.item(i, "values")[0] for i in ids]
                assert got == ["☑"] * len(ids), f"全选后有没勾上的：{got}"
                txt = sum_text(tw)
                assert f"已选 {len(ids)} 人" in txt, f"全选后汇总行不对：{txt}"
                # 候选不在名称库里 → 必须提醒（认错了客户发言会被当成我方答复）
                assert "⚠" in txt and "可疑客户" in txt, \
                    f"勾上扫描候选却没提醒：{txt}"

                find_widgets(tw, tk.ttk.Button, "全不选")[0].invoke()
                a.update()
                got = [tvw.item(i, "values")[0] for i in ids]
                assert got == ["☐"] * len(ids), f"全不选后还有勾着的：{got}"
                txt = sum_text(tw)
                assert "已选 0 人" in txt, f"全不选后汇总行不对：{txt}"
            finally:
                tw.destroy()
                a.update()
        finally:
            a.self_candidates = []

    check("勾选：「全选 / 全不选」能勾上或清空整表，汇总行跟着变",
          select_all_and_none)

    # ---- 扫描：这条路径曾经整条失效（轮询循环没被启动）----
    # 症状是点「开始扫描」后永远停在"正在准备…"，后台其实抓到了数据，
    # 只是没人读结果队列。所以这里必须断言状态**真的会变**，不能只看按钮在不在。

    def scan_actually_finishes():
        import time as _time

        a.contacts = [Contact(cid="c1", name="某交付群", kind="group", msg_count=10)]
        a.picked = {"c1"}
        a.cfg["source"] = "demo"
        a.cfg["source_path"] = ""
        a.var_start.set("2026-01-01")
        a.var_end.set("2026-12-31")
        a.self_candidates = []
        a._open_self_picker()
        a.update()
        t2 = picker_window()
        btn = find_widgets(t2, tk.ttk.Button, "开始扫描")
        assert btn, "没有找到开始扫描按钮"
        btn[0].invoke()
        a.update()

        # 泵事件循环，等状态离开「正在准备…」
        texts = []
        for _ in range(150):                       # 最多 15 秒
            a.update()
            _time.sleep(0.1)
            txt = _state_text(t2)
            if txt and txt not in texts:
                texts.append(txt)
            if txt and "正在准备" not in txt and "正在看" not in txt:
                break
        txt = _state_text(t2)
        assert txt, f"状态栏一直没更新（读到的状态历史：{texts}）"
        assert "正在准备" not in txt, f"卡在「正在准备…」没动：状态={txt!r}，历史={texts}"
        assert btn[0].cget("state") != "disabled", "扫描结束后按钮该恢复可点"

    def _state_text(win):
        """把选择窗口里那个状态标签的文字读出来。"""
        for w in find_widgets(win, tk.ttk.Label):
            tv = str(w.cget("textvariable") or "")
            if tv:
                try:
                    v = w.getvar(tv)
                except Exception:                                    # noqa: BLE001
                    continue
                if v and ("准备" in v or "找到" in v or "扫描" in v or "勾选" in v
                          or "正在看" in v):
                    return v
        return ""

    check("开始扫描：状态会真的更新（不会卡在「正在准备…」）", scan_actually_finishes)

    def add_button_works():
        a.self_candidates = []
        a._open_self_picker()
        a.update()
        t3 = picker_window()
        ent = find_widgets(t3, tk.ttk.Entry)
        assert ent, "没有找到手动输入框"
        ent[0].insert(0, "按钮加的同事")
        btn = find_widgets(t3, tk.ttk.Button, "添加")
        assert btn, "没有找到「添加」按钮"
        btn[0].invoke()
        a.update()
        tv3, rows3 = rows_of(t3)
        d = {r[1]: r for r in rows3}
        assert "按钮加的同事" in d, f"点「添加」没把人加进去：{list(d)}"
        assert d["按钮加的同事"][0] == "☑", f"添加后应勾上：{d['按钮加的同事']}"
        # 点「确定」应把它存进名称库
        find_widgets(t3, tk.ttk.Button, "确定")[0].invoke()
        a.update()
        assert "按钮加的同事" in (a.cfg.get("self_name_library") or []), \
            f"没入库：{a.cfg.get('self_name_library')}"

    check("名称库：「添加」按钮能加人并入库", add_button_works)

    def export_template():
        import tempfile

        from openpyxl import load_workbook

        out = Path(tempfile.mkdtemp(prefix="wxcsm_tpl_")) / "模版.xlsx"
        real_save = appmod.filedialog.asksaveasfilename
        appmod.filedialog.asksaveasfilename = lambda *x, **k: str(out)
        try:
            a.self_candidates = []
            a._open_self_picker()
            a.update()
            t4 = picker_window()
            # 这次没扫描（self_candidates 清空了），所以表里全是名称库成员——
            # 模版带出的应当**正好**是这些人，一个不多一个不少。
            expect = {r[1] for r in rows_of(t4)[1]}
            assert expect, "名称库是空的，没法验证模版带库"
            btn = find_widgets(t4, tk.ttk.Button, "导出 Excel 模版")
            assert btn, "没有找到「导出 Excel 模版」按钮"
            btn[0].invoke()
            a.update()
        finally:
            appmod.filedialog.asksaveasfilename = real_save

        assert out.exists(), f"模版没写出来：{out}"
        wb = load_workbook(out)
        ws = wb.active
        heads = [c.value for c in ws[1]]
        print(f"        模版表头：{heads}")
        assert any("姓名" in str(h) for h in heads), f"模版没有姓名列：{heads}"
        # 必填项的表头要标红
        import core.name_library as NL

        req = [c for c in ws[1] if "必填" in str(c.value)]
        assert req, f"模版里没有标出必填项：{heads}"
        for c in req:
            col = c.font.color
            rgb = getattr(col, "rgb", None) if col is not None else None
            assert rgb and rgb.endswith(NL.TEMPLATE_REQUIRED_COLOR[-6:]), \
                f"必填项表头没标红：{c.value} font={rgb}"
        # 可选列不该标红
        for c in ws[1]:
            if "可选" in str(c.value):
                col = c.font.color
                rgb = getattr(col, "rgb", None) if col is not None else None
                assert not (rgb and rgb.endswith(NL.TEMPLATE_REQUIRED_COLOR[-6:])), \
                    f"可选列被标红了：{c.value}"
        wb.close()

        # 模版要带出当前名称库的人——用户就是在这张表上改（加人/删行）再导入回来。
        cols = NL.read_columns(out)
        names = [n for c in cols for n in c.names]
        assert set(names) == expect, f"模版带出的库成员不对：{names} ≠ {expect}"
        # 表头不能被当成一个人带出去（名称库是子串匹配判我方身份的，泛词会误伤客户发言）
        assert not any("姓名" in n or "必填" in n or "可选" in n for n in names), \
            f"表头被当成人名带出去了：{names}"
        print(f"        模版带出的库成员：{names}")

    check("名称库：导出 Excel 模版（必填表头标红、带出名称库、表头不算人名）",
          export_template)

    print("【7】导出（真的写文件）")
    # 写到临时目录：测试不该往仓库根目录丢文件
    import tempfile
    out = Path(tempfile.mkdtemp(prefix="wxcsm_gui_smoke_")) / "冒烟.xlsx"
    if out.exists():
        out.unlink()
    a.var_export.set(str(out))
    a.var_mode.set("overwrite")

    def do_export_both():
        a.var_format.set("both")
        a._update_format_ui()
        a.update()
        a._export()
        a.update()
        from openpyxl import load_workbook
        assert out.exists(), "导出后文件不存在"
        wb = load_workbook(out)
        assert "客户沟通总结" in wb.sheetnames, wb.sheetnames
        assert "客户问题答复清单" in wb.sheetnames, wb.sheetnames
        ws = wb["客户问题答复清单"]
        n = ws.max_row - 1
        assert n == 3, f"问答清单应有 3 行，实际 {n}"
        assert ws.cell(1, 4).value == "客户问题", ws.cell(1, 4).value
        ws2 = wb["客户沟通总结"]
        assert ws2.max_row - 1 == 1, f"总结应有 1 行，实际 {ws2.max_row - 1}"

    check("「两者都要」导出 → 两张 sheet 都在且行数正确", do_export_both)

    def qa_only_keeps_summary_sheet():
        # 只导出问答时，已有的总结 sheet 不能被删掉（历史上真踩过这个 bug）
        a.var_format.set("qa")
        a._update_format_ui()
        a.update()
        a._export()
        a.update()
        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert "客户沟通总结" in wb.sheetnames, f"总结 sheet 被删了：{wb.sheetnames}"
        assert "客户问题答复清单" in wb.sheetnames

    check("「仅问答」导出不会删掉总结 sheet", qa_only_keeps_summary_sheet)

    print("【8】清空结果")
    a._clear_res()
    a.update()
    check("清空后总结与问答都为空",
          lambda: not a.summaries and not a.qa_items)

    print("【9】真实跑一遍 demo 数据（串起 _on_result 全链路）")

    def real_run():
        from core import pipeline as P
        cfg = dict(a.cfg)
        cfg.update(source="demo", source_path="", output_format="both",
                   engine="offline", min_chars=50, max_chars=200,
                   export_path=str(out), export_mode="overwrite")
        contacts = P.load_contacts(cfg)
        assert contacts, "demo 源没有联系人"
        start, end = P.preset_range("30d")
        a.summaries.clear()
        a.qa_items.clear()
        res = P.run(cfg, contacts[:2], start, end, do_export=True)
        a._on_result(res)
        a.update()
        assert a.summaries, "走完 _on_result 后没有总结"
        assert a.qa_items, "走完 _on_result 后没有问答条目"
        assert len(a.tv_res.get_children()) == len(a.qa_items), \
            f"表格行数 {len(a.tv_res.get_children())} != 问答条目 {len(a.qa_items)}"
        assert "问答" in a.lbl_res.cget("text") or "条" in a.lbl_res.cget("text")

    check("demo 数据跑通并渲染到结果区", real_run)

    def real_export():
        from openpyxl import load_workbook
        a.var_format.set("both")
        a._update_format_ui()
        a.update()
        a._export()
        a.update()
        wb = load_workbook(out)
        assert "客户沟通总结" in wb.sheetnames and "客户问题答复清单" in wb.sheetnames
        assert wb["客户问题答复清单"].max_row - 1 == len(a.qa_items)

    check("真实结果导出两张 sheet 行数一致", real_export)

    a.destroy()
    return 1 if FAILS else 0


if __name__ == "__main__":
    try:
        code = main()
    finally:
        # 还原真实配置：_sync_cfg 会写盘，测试值不能留在用户配置里
        if BACKUP is not None:
            CFG.write_text(BACKUP, encoding="utf-8")
            print("\n（已还原 config.json）")
    print()
    if FAILS:
        print(f"失败 {len(FAILS)} 项：" + "、".join(FAILS))
    else:
        print("GUI 冒烟测试全部通过")
    sys.exit(code)
